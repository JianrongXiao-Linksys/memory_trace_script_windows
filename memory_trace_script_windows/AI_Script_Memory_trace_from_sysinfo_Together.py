#!/usr/bin/env python
# -*- coding: utf-8 -*-

# 20251212 - 最終解決方案：生成於同級目錄，再和txt一同移動。

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.chart import LineChart, BarChart, Reference
from datetime import datetime
from collections import defaultdict
import os
import re
import shutil 

# --- 關鍵字設定 ---
MEMORY_KEYWORDS = {
    "MemAvailable": "MemAvailable:",
    "AnonPages": "AnonPages:", 
    "SUnreclaim": "SUnreclaim:", 
}
# -----------------------------

# --- 數據讀取與處理函數 (與前一版本相同，略過) ---
def get_all_log_files(target_dir="."):
    """查找當前或指定目錄中所有 sysinfo .txt 檔案並按名稱排序"""
    all_files = os.listdir(target_dir)
    log_files = [f for f in all_files if str(f).endswith(".txt") and "sysinfo" in f]
    log_files.sort()
    return log_files

def extract_timestamp_from_filename(filename, format_type='full'):
    """從檔名中提取時間戳記或日期。"""
    try:
        match = re.search(r'_(\d{4}-\d{2}-\d{2})_(\d{6})\.txt$', filename)
        if match:
            date_str = match.group(1)
            time_str = match.group(2)
            if format_type == 'date':
                return date_str
            dt_obj = datetime.strptime(f"{date_str}_{time_str}", "%Y-%m-%d_%H%M%S")
            return dt_obj.strftime("%Y-%m-%d %H:%M:%S") 
        return filename
    except Exception:
        return filename

def parse_memory_value(line, keyword):
    """從包含關鍵字的行中提取記憶體數值 (以 KB 為單位)"""
    try:
        match = re.search(r'\s+(\d+)\s*(?:k|K|kb|KB)?\b', line[len(keyword):], re.IGNORECASE)
        if match:
            return int(match.group(1))
        return 0
    except:
        return 0

def get_memory_info(filename):
    """從單個 sysinfo 檔案中提取所有關鍵字對應的記憶體值。(假設檔案在 CWD)"""
    memory_values = {key: 0 for key in MEMORY_KEYWORDS}
    try:
        with open(filename, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.readlines()
        for line in content:
            for key, keyword in MEMORY_KEYWORDS.items():
                if keyword in line and memory_values[key] == 0:
                    memory_values[key] = parse_memory_value(line, keyword)
    except Exception as e:
        pass
    if any(memory_values.values()):
        return memory_values
    return None

def collect_all_data(all_log_files):
    """讀取所有檔案並收集每筆記錄的數據點。"""
    all_memory_data = {}
    for each_file in all_log_files:
        full_time_str = extract_timestamp_from_filename(each_file, format_type='full')
        memory_values = get_memory_info(each_file)
        if memory_values and full_time_str != each_file:
            all_memory_data[full_time_str] = memory_values
    return all_memory_data

def calculate_daily_average(all_memory_data):
    """計算每日的記憶體指標平均值。"""
    daily_sums_and_counts = defaultdict(lambda: defaultdict(lambda: {'sum': 0, 'count': 0}))
    daily_average_data = {}
    metrics = list(MEMORY_KEYWORDS.keys())

    for full_time_str, data in all_memory_data.items():
        date_str = full_time_str.split(' ')[0] 
        for metric in metrics:
            value = data.get(metric, 0)
            if value > 0:
                daily_sums_and_counts[date_str][metric]['sum'] += value
                daily_sums_and_counts[date_str][metric]['count'] += 1
                
    for date, metrics_data in daily_sums_and_counts.items():
        daily_average_data[date] = {}
        for metric in metrics:
            s = metrics_data[metric]['sum']
            c = metrics_data[metric]['count']
            average = int(s / c) if c > 0 else 0
            daily_average_data[date][metric] = average
            
    return daily_average_data


# --- Excel 寫入函數 ---

def write_daily_average_sheet(worksheet, daily_average_data):
    """寫入每日平均數據和長條圖，並確保 Y 軸從 0 開始。"""
    metrics = list(MEMORY_KEYWORDS.keys())
    headers = ["Date"] + [f"{m} Average (KB)" for m in metrics]
    
    for col_idx, header in enumerate(headers, 1):
        col_letter = chr(ord('A') + col_idx - 1)
        worksheet["{}1".format(col_letter)] = header
        worksheet.column_dimensions[col_letter].width = 25
        worksheet["{}1".format(col_letter)].alignment = Alignment(horizontal='center')
        
    worksheet.freeze_panes = "A2"
    sorted_dates = sorted(daily_average_data.keys())
    start_row = 1
    
    for date_str in sorted_dates:
        start_row += 1
        data = daily_average_data[date_str]
        date_cell = "A{}".format(start_row)
        worksheet[date_cell] = date_str
        worksheet[date_cell].alignment = Alignment(horizontal='center')
        for i, metric_name in enumerate(metrics):
            col_letter = chr(ord('B') + i)
            data_cell = f"{col_letter}{start_row}"
            worksheet[data_cell] = data.get(metric_name, 0)
            worksheet[data_cell].alignment = Alignment(horizontal='center')
             
    metric_cols = list(range(2, 2 + len(metrics)))
    for i, metric_name in enumerate(metrics):
        col_index = metric_cols[i]
        chart = BarChart()
        chart.type = "col" 
        chart.style = 10
        chart.title = f"Daily Average of {metric_name}"
        chart.x_axis.title = "Date"
        chart.y_axis.title = f"{metric_name} Average (KB)"
        
        # 確保 Y 軸最小值為 0
        chart.y_axis.scaling.min = 0 
        
        data_ref = Reference(worksheet, min_col=col_index, min_row=1, max_col=col_index, max_row=start_row)
        categories_ref = Reference(worksheet, min_col=1, min_row=2, max_row=start_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(categories_ref)
        row_offset = i * 15
        chart_cell = f"E{2 + row_offset}"
        worksheet.add_chart(chart, chart_cell)

def write_time_series_sheet(worksheet, all_memory_data):
    """寫入每筆記錄數據和折線圖，並確保 Y 軸從 0 開始。"""
    metrics = list(MEMORY_KEYWORDS.keys())
    headers = ["Date & Time"] + [f"{m} (KB)" for m in metrics]
    
    for col_idx, header in enumerate(headers, 1):
        col_letter = chr(ord('A') + col_idx - 1)
        worksheet["{}1".format(col_letter)] = header
        worksheet.column_dimensions[col_letter].width = 22
        worksheet["{}1".format(col_letter)].alignment = Alignment(horizontal='center')
        
    worksheet.freeze_panes = "A2"
    sorted_times = sorted(all_memory_data.keys())
    start_row = 1
    
    for time_str in sorted_times:
        start_row += 1
        data = all_memory_data[time_str]
        time_cell = "A{}".format(start_row)
        worksheet[time_cell] = time_str
        worksheet[time_cell].alignment = Alignment(horizontal='center')
        for i, metric_name in enumerate(metrics):
            col_letter = chr(ord('B') + i)
            data_cell = f"{col_letter}{start_row}"
            worksheet[data_cell] = data.get(metric_name, 0)
            worksheet[data_cell].alignment = Alignment(horizontal='center')

    metric_cols = list(range(2, 2 + len(metrics)))
    for i, metric_name in enumerate(metrics):
        col_index = metric_cols[i]
        chart = LineChart()
        chart.title = f"Time Series of {metric_name}"
        chart.x_axis.title = "Date & Time"
        chart.y_axis.title = f"{metric_name} (KB)"
        
        # 確保 Y 軸最小值為 0
        chart.y_axis.scaling.min = 0
        
        data_ref = Reference(worksheet, min_col=col_index, min_row=1, max_col=col_index, max_row=start_row)
        categories_ref = Reference(worksheet, min_col=1, min_row=2, max_row=start_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(categories_ref)
        row_offset = i * 15 
        chart_cell = f"E{2 + row_offset}"
        worksheet.add_chart(chart, chart_cell)

def create_combined_excel(all_memory_data, daily_average_data, output_filepath_temp):
    """創建單一 Excel 檔案，儲存到同級目錄的臨時路徑。"""
    workbook = Workbook()
    
    # 1. 寫入每日平均分頁 
    ws_avg = workbook.active
    ws_avg.title = "DailyAverage"
    write_daily_average_sheet(ws_avg, daily_average_data)
    
    # 2. 寫入每筆記錄分頁
    ws_ts = workbook.create_sheet(title="TimeSeries")
    write_time_series_sheet(ws_ts, all_memory_data)

    # 確保 DailyAverage 仍是 active sheet (預設分頁)
    workbook.active = 0 

    # *** 儲存到臨時的、同級目錄的絕對路徑 ***
    workbook.save(output_filepath_temp)
    print(f"Combined Excel file saved successfully to current directory: {output_filepath_temp}")


# --- 主程式碼 ---
if __name__ == "__main__":    
    
    # 1. 設置路徑變數
    current_dir_absolute = os.path.abspath(os.getcwd())
    log_root_name = os.path.basename(current_dir_absolute)
    current_time_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    # 新的輸出資料夾名稱 
    output_folder_name = "MemoryAnalysis_{}".format(current_time_str)
    
    # 新的輸出資料夾絕對路徑 (在當前目錄下)
    output_folder_path_absolute = os.path.join(current_dir_absolute, output_folder_name)
    
    # 最終 Excel 檔案名稱
    filename_base = "Memory_Analysis_Combined_{}.xlsx".format(log_root_name)
    
    # *** 臨時 Excel 儲存路徑：同級目錄 + 檔名 ***
    filename_combined_path_temp = os.path.join(current_dir_absolute, filename_base)
    
    # 最終 Excel 歸檔路徑：新資料夾 + 檔名
    filename_combined_path_final = os.path.join(output_folder_path_absolute, filename_base)
    
    # 2. 查找檔案 (在當前目錄)
    all_log_files = get_all_log_files() 
    if len(all_log_files) == 0:
        print("No sysinfo (.txt) file found in the current directory.")
        print("Please ensure the script is executed in the folder containing the log files (e.g., the 'sysinfo' folder).")
        exit(10)
        
    # 3. 創建輸出資料夾
    try:
        # 使用絕對路徑創建資料夾
        os.makedirs(output_folder_path_absolute, exist_ok=True)
        print(f"Created output folder: {output_folder_name}")
    except Exception as e:
        print(f"Error creating folder: {e}. Aborting.")
        exit(1)
        
    # 4. 數據收集與處理
    print("Collecting and processing data...")
    all_memory_data = collect_all_data(all_log_files) 

    if len(all_memory_data) == 0:
        print("No memory data extracted. Please check the content of sysinfo files.")
    else:
        # 計算每日平均
        daily_average_data = calculate_daily_average(all_memory_data)
        
        # 5. 生成整合 Excel 檔案 (儲存到同級目錄)
        print("Generating combined Excel file to the current directory...")
        try:
            create_combined_excel(all_memory_data, daily_average_data, filename_combined_path_temp)
        except Exception as e:
            print(f"CRITICAL ERROR: Failed to create Excel file at current directory: {filename_combined_path_temp}. Error: {e}")
            
    # --- 6 & 7. 檔案歸檔 (移動 Excel 和 Log) ---
    print("\n--- Starting Archiving: Moving files to output folder ---")

    # 6. 移動 Excel 檔案
    if os.path.exists(filename_combined_path_temp):
        print(f"Moving Excel file to final folder: {output_folder_name}...")
        try:
            # 從同級目錄移動到最終歸檔路徑
            shutil.move(filename_combined_path_temp, filename_combined_path_final)
            print("Excel file moved successfully.")
        except Exception as e:
            print(f"ERROR: Failed to move Excel file: {e}. Please check if the file is open or if permissions are restricted.")
    else:
        print("WARNING: Excel file was not found in the current directory, skipping Excel move operation.")
            
    # 7. 移動原始檔案
    print("Moving original sysinfo files...")
    
    log_files_to_move = get_all_log_files() 
    
    try:
        for file_name in log_files_to_move:
            source_log_path_absolute = os.path.join(current_dir_absolute, file_name)
            destination_path = os.path.join(output_folder_path_absolute, file_name)
            
            # 使用 shutil.move 進行移動
            if os.path.isfile(source_log_path_absolute):
                 shutil.move(source_log_path_absolute, destination_path)
        print("Log files moving completed.")
    except Exception as e:
        print(f"Error moving log files: {e}. Please manually move the files if necessary.")
        
    print(f"\nAnalysis completed! Results and original files are in the folder: {output_folder_path_absolute}")
